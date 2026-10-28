
import sqlite3
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

import streamlit as st
from docx import Document
from docx.shared import Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable
)


# ==================================================
# 1. CONFIGURATION
# ==================================================

st.set_page_config(
    page_title="CIAT DocIA",
    page_icon="📄",
    layout="wide"
)


# ==================================================
# 2. DESIGN
# ==================================================

st.markdown("""
<style>

.stApp {
    background-color: #f5f7f6;
}

[data-testid="stSidebar"] {
    background-color: #14532d;
}

[data-testid="stSidebar"] * {
    color: white;
}

h1, h2, h3 {
    color: #166534;
}

.stButton > button {
    background-color: #15803d;
    color: white;
    border-radius: 8px;
    border: none;
    padding: 10px 20px;
}

.stButton > button:hover {
    background-color: #166534;
    color: white;
}

div[data-testid="stDownloadButton"] button {
    background-color: #166534;
    color: white;
    border-radius: 8px;
}

</style>
""", unsafe_allow_html=True)


# ==================================================
# 3. BASE DE DONNEES SQLITE
# ==================================================

def initialiser_base():

    connexion = sqlite3.connect("ciat_docia.db")

    curseur = connexion.cursor()

    curseur.execute("""
        CREATE TABLE IF NOT EXISTS fiches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            titre TEXT NOT NULL,
            responsable TEXT NOT NULL,
            contenu TEXT NOT NULL,
            date_creation TEXT NOT NULL
        )
    """)

    connexion.commit()

    connexion.close()


initialiser_base()


# ==================================================
# 4. ENREGISTRER UNE FICHE
# ==================================================

def enregistrer_fiche(titre, responsable, contenu, date_creation):

    connexion = sqlite3.connect("ciat_docia.db")

    curseur = connexion.cursor()

    curseur.execute("""
        INSERT INTO fiches (
            titre,
            responsable,
            contenu,
            date_creation
        )
        VALUES (?, ?, ?, ?)
    """, (
        titre,
        responsable,
        contenu,
        date_creation
    ))

    connexion.commit()

    connexion.close()


# ==================================================
# 5. RECUPERER LES FICHES
# ==================================================

def recuperer_fiches():

    connexion = sqlite3.connect("ciat_docia.db")

    curseur = connexion.cursor()

    curseur.execute("""
        SELECT id, titre, responsable, contenu, date_creation
        FROM fiches
        ORDER BY id DESC
    """)

    fiches = curseur.fetchall()

    connexion.close()

    return fiches


# ==================================================
# 6. GENERER LE DOCUMENT WORD
# ==================================================

def creer_word(titre, responsable, contenu, date_creation):

    document = Document()

    # Titre principal
    titre_document = document.add_heading(
        "CIAT DocIA",
        level=0
    )

    titre_document.alignment = 1

    # Sous-titre
    sous_titre = document.add_paragraph()

    sous_titre.alignment = 1

    texte = sous_titre.add_run(
        "FICHE DE PROJET"
    )

    texte.bold = True

    texte.font.size = Pt(16)

    texte.font.color.rgb = RGBColor(22, 101, 52)

    document.add_paragraph()

    # Informations générales
    document.add_heading(
        "Informations générales",
        level=1
    )

    document.add_paragraph(
        f"Titre du projet : {titre}"
    )

    document.add_paragraph(
        f"Responsable : {responsable}"
    )

    document.add_paragraph(
        f"Date de création : {date_creation}"
    )

    # Contenu
    document.add_heading(
        "Description du projet",
        level=1
    )

    lignes = contenu.split("\n")

    for ligne in lignes:

        ligne = ligne.strip()

        if not ligne:
            continue

        if ligne.startswith("# "):

            document.add_heading(
                ligne[2:],
                level=1
            )

        elif ligne.startswith("## "):

            document.add_heading(
                ligne[3:],
                level=2
            )

        elif ligne.startswith("- "):

            document.add_paragraph(
                ligne[2:],
                style="List Bullet"
            )

        else:

            document.add_paragraph(ligne)

    # Pied de document
    document.add_paragraph()

    pied = document.add_paragraph(
        "Document généré par CIAT DocIA."
    )

    pied.alignment = 1

    # Enregistrer dans la mémoire temporaire
    fichier = BytesIO()

    document.save(fichier)

    fichier.seek(0)

    return fichier.getvalue()


# ==================================================
# 7. GENERER LE DOCUMENT PDF
# ==================================================

def creer_pdf(titre, responsable, contenu, date_creation):

    fichier = BytesIO()

    document = SimpleDocTemplate(
        fichier,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm
    )

    styles = getSampleStyleSheet()

    style_titre = ParagraphStyle(
        name="TitreCIAT",
        parent=styles["Title"],
        fontSize=22,
        leading=28,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#166534"),
        spaceAfter=15
    )

    style_section = ParagraphStyle(
        name="SectionCIAT",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#166534"),
        spaceBefore=12,
        spaceAfter=6
    )

    style_normal = ParagraphStyle(
        name="NormalCIAT",
        parent=styles["BodyText"],
        fontSize=10,
        leading=15,
        spaceAfter=6
    )

    elements = []

    # Titre
    elements.append(
        Paragraph(
            "CIAT DocIA",
            style_titre
        )
    )

    elements.append(
        Paragraph(
            "FICHE DE PROJET",
            style_section
        )
    )

    elements.append(
        HRFlowable(
            width="100%",
            thickness=1,
            color=colors.HexColor("#166534")
        )
    )

    elements.append(
        Spacer(1, 12)
    )

    # Informations générales
    elements.append(
        Paragraph(
            "<b>Informations générales</b>",
            style_section
        )
    )

    elements.append(
        Paragraph(
            "<b>Titre :</b> " + escape(titre),
            style_normal
        )
    )

    elements.append(
        Paragraph(
            "<b>Responsable :</b> " + escape(responsable),
            style_normal
        )
    )

    elements.append(
        Paragraph(
            "<b>Date de création :</b> " + escape(date_creation),
            style_normal
        )
    )

    elements.append(
        Spacer(1, 10)
    )

    # Contenu
    for ligne in contenu.split("\n"):

        ligne = ligne.strip()

        if not ligne:
            elements.append(
                Spacer(1, 5)
            )

            continue

        if ligne.startswith("# "):

            elements.append(
                Paragraph(
                    escape(ligne[2:]),
                    style_section
                )
            )

        elif ligne.startswith("## "):

            elements.append(
                Paragraph(
                    escape(ligne[3:]),
                    style_section
                )
            )

        elif ligne.startswith("- "):

            elements.append(
                Paragraph(
                    "• " + escape(ligne[2:]),
                    style_normal
                )
            )

        else:

            elements.append(
                Paragraph(
                    escape(ligne),
                    style_normal
                )
            )

    document.build(elements)

    fichier.seek(0)

    return fichier.getvalue()


# ==================================================
# 8. BARRE LATERALE
# ==================================================

with st.sidebar:

    st.title("📄 CIAT DocIA")

    st.write(
        "Générateur de documents"
    )

    st.divider()

    page = st.radio(
        "Navigation",
        [
            "🏠 Accueil",
            "📝 Créer une fiche",
            "📚 Historique"
        ]
    )

    st.divider()

    st.caption(
        "Prototype de démonstration"
    )


# ==================================================
# 9. PAGE ACCUEIL
# ==================================================

if page == "🏠 Accueil":

    st.title("📄 Bienvenue dans CIAT DocIA")

    st.write(
        "Application de création et de gestion "
        "de fiches de projets."
    )

    st.success(
        "Bienvenue dans votre espace de travail."
    )

    st.subheader("Fonctionnalités")

    col1, col2, col3 = st.columns(3)

    with col1:

        st.info(
            "📝 Créer une fiche"
        )

        st.write(
            "Remplir un formulaire pour "
            "préparer une fiche de projet."
        )

    with col2:

        st.info(
            "📚 Historique"
        )

        st.write(
            "Consulter les fiches enregistrées."
        )

    with col3:

        st.info(
            "📥 Export"
        )

        st.write(
            "Télécharger les fiches en PDF ou Word."
        )

    st.divider()

    st.subheader("État de la démonstration")

    st.write(
        "Les fiches sont enregistrées dans une "
        "base de données SQLite locale."
    )

    st.warning(
        "Cette version utilise un modèle de document. "
        "La génération réelle par intelligence artificielle "
        "n'est pas encore intégrée."
    )


# ==================================================
# 10. PAGE CREATION
# ==================================================

elif page == "📝 Créer une fiche":

    st.title("📝 Créer une fiche de projet")

    st.write(
        "Remplissez les informations ci-dessous."
    )

    with st.form("formulaire_fiche"):

        titre = st.text_input(
            "Titre du projet",
            placeholder="Exemple : Projet agricole"
        )

        responsable = st.text_input(
            "Responsable du projet",
            placeholder="Nom du responsable"
        )

        probleme = st.text_area(
            "Problématique",
            placeholder="Décrivez le problème à résoudre"
        )

        solution = st.text_area(
            "Solution proposée",
            placeholder="Décrivez la solution envisagée"
        )

        objectifs = st.text_area(
            "Objectifs",
            placeholder="Indiquez les objectifs du projet"
        )

        resultat = st.text_area(
            "Résultats attendus",
            placeholder="Indiquez les résultats attendus"
        )

        bouton_creer = st.form_submit_button(
            "💾 Créer la fiche"
        )

    if bouton_creer:

        if not titre.strip():

            st.error(
                "Veuillez renseigner le titre du projet."
            )

        elif not responsable.strip():

            st.error(
                "Veuillez renseigner le responsable."
            )

        elif not probleme.strip():

            st.error(
                "Veuillez renseigner la problématique."
            )

        elif not solution.strip():

            st.error(
                "Veuillez renseigner la solution proposée."
            )

        elif not objectifs.strip():

            st.error(
                "Veuillez renseigner les objectifs."
            )

        else:

            date_creation = datetime.now().strftime(
                "%d/%m/%Y %H:%M"
            )

            contenu = f"""
# 1. Problématique

{probleme}

# 2. Solution proposée

{solution}

# 3. Objectifs

{objectifs}

# 4. Résultats attendus

{resultat if resultat.strip() else "À préciser."}

# 5. Informations à préciser

- Calendrier du projet : à préciser
- Budget : à préciser
- Ressources nécessaires : à préciser
- Indicateurs de performance : à préciser

# 6. État de la génération

Document préparé à partir d'un modèle local de démonstration.
Le contenu n'a pas été généré par une intelligence artificielle.
"""

            enregistrer_fiche(
                titre,
                responsable,
                contenu,
                date_creation
            )

            st.session_state["fiche_creee"] = {
                "titre": titre,
                "responsable": responsable,
                "contenu": contenu,
                "date_creation": date_creation
            }

            st.success(
                "La fiche a été enregistrée dans SQLite."
            )

    # Afficher la fiche créée
    if "fiche_creee" in st.session_state:

        fiche = st.session_state["fiche_creee"]

        st.divider()

        st.subheader("📄 Aperçu de la fiche")

        st.write(
            "**Titre :**",
            fiche["titre"]
        )

        st.write(
            "**Responsable :**",
            fiche["responsable"]
        )

        st.write(
            "**Date :**",
            fiche["date_creation"]
        )

        st.markdown(
            fiche["contenu"]
        )

        # Télécharger Word
        word = creer_word(
            fiche["titre"],
            fiche["responsable"],
            fiche["contenu"],
            fiche["date_creation"]
        )

        st.download_button(
            label="📄 Télécharger la fiche en Word",
            data=word,
            file_name="fiche_CIAT_DocIA.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key="telecharger_word_creation"
        )

        # Télécharger PDF
        pdf = creer_pdf(
            fiche["titre"],
            fiche["responsable"],
            fiche["contenu"],
            fiche["date_creation"]
        )

        st.download_button(
            label="📥 Télécharger la fiche en PDF",
            data=pdf,
            file_name="fiche_CIAT_DocIA.pdf",
            mime="application/pdf",
            key="telecharger_pdf_creation"
        )


# ==================================================
# 11. PAGE HISTORIQUE
# ==================================================

elif page == "📚 Historique":

    st.title("📚 Historique des fiches")

    fiches = recuperer_fiches()

    if not fiches:

        st.info(
            "Aucune fiche enregistrée pour le moment."
        )

        st.write(
            "Créez une fiche pour la retrouver ici."
        )

    else:

        st.success(
            f"{len(fiches)} fiche(s) enregistrée(s)."
        )

        for fiche in fiches:

            identifiant = fiche[0]

            titre = fiche[1]

            responsable = fiche[2]

            contenu = fiche[3]

            date_creation = fiche[4]

            with st.expander(
                f"📄 {titre} — {date_creation}"
            ):

                st.write(
                    "**Responsable :**",
                    responsable
                )

                st.markdown(
                    contenu
                )

                # Word
                word = creer_word(
                    titre,
                    responsable,
                    contenu,
                    date_creation
                )

                st.download_button(
                    label="📄 Télécharger en Word",
                    data=word,
                    file_name=f"fiche_CIAT_DocIA_{identifiant}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key=f"word_{identifiant}"
                )

                # PDF
                pdf = creer_pdf(
                    titre,
                    responsable,
                    contenu,
                    date_creation
                )

                st.download_button(
                    label="📥 Télécharger en PDF",
                    data=pdf,
                    file_name=f"fiche_CIAT_DocIA_{identifiant}.pdf",
                    mime="application/pdf",
                    key=f"pdf_{identifiant}"
                )


# ==================================================
# 12. FIN DE L'APPLICATION
# ==================================================

st.divider()

st.caption(
    "CIAT DocIA — Prototype de démonstration"
)
